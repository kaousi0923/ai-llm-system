import os
import shutil
import re
from typing import List, Tuple

from fastapi import UploadFile, HTTPException

from langchain.docstore.document import Document

import pdfplumber
from docx import Document as DocxDocument


def extract_tags(name: str) -> Tuple[str, List[str]]:
    match = re.search(r'\((.*?)\)', name)
    if match:
        tags = [tag.strip() for tag in match.group(1).split(',')]
        clean_name = re.sub(r'\(.*?\)', '', name).strip()
    else:
        tags = []
        clean_name = name
    return clean_name, tags


def extract_text_from_pdf(path: str) -> str:
    texts = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            # 擷取純文字
            text = page.extract_text() or ''
            texts.append(text)

            # 擷取表格
            tables = page.extract_tables()
            for table in tables:
                for row in table:
                    # 將 None 轉為空字串，避免 join 時出錯
                    safe_row = [cell if cell is not None else '' for cell in row]
                    row_text = "\t".join(safe_row)
                    texts.append(row_text)
    return "\n".join(texts)


def extract_text_from_docx(path: str) -> str:
    docx = DocxDocument(path)
    output = []

    # 1. 擷取所有段落文字（包含標題與敘述性文字）
    for para in docx.paragraphs:
        text = para.text.strip()
        if text:
            output.append(text)

    # 2. 擷取所有表格內容（轉為 Markdown 樣式）
    for table_idx, table in enumerate(docx.tables):
        rows = table.rows
        if not rows or len(rows) < 1:
            continue  # 略過空表格

        # 嘗試判斷是否第一列為標題
        first_row = [cell.text.strip() for cell in rows[0].cells]
        has_header = all(col != "" for col in first_row)

        output.append(f"\n【表格 {table_idx + 1}】")
        if has_header:
            # markdown 表頭
            output.append("| " + " | ".join(first_row) + " |")
            output.append("| " + " | ".join(["---"] * len(first_row)) + " |")
            data_rows = rows[1:]
        else:
            data_rows = rows

        for row in data_rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):  # 避免整列皆空白
                output.append("| " + " | ".join(cells) + " |")

    return "\n".join(output)


def process_uploaded_files(files: List[UploadFile], temp_dir: str) -> Tuple[List[Document], List[str]]:
    os.makedirs(temp_dir, exist_ok=True)
    documents = []
    file_names = []

    for file in files:
        file_path = os.path.join(temp_dir, file.filename)
        with open(file_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)

        try:
            if file.filename.endswith('.pdf'):
                try:
                    full_text = extract_text_from_pdf(file_path)
                    docs = [Document(page_content=full_text, metadata={"source": file.filename})]
                except Exception as e:
                    raise RuntimeError(f"pdfplumber 讀取失敗：{file.filename}，錯誤：{str(e)}")

            elif file.filename.endswith('.docx'):
                try:
                    full_text = extract_text_from_docx(file_path)
                    docs = [Document(page_content=full_text, metadata={"source": file.filename})]
                except Exception as e:
                    raise RuntimeError(f"docx 讀取失敗：{file.filename}，錯誤：{str(e)}")

            elif file.filename.endswith('.txt'):
                try:
                    with open(file_path, "r", encoding="utf-8") as f:
                        text = f.read()
                except UnicodeDecodeError:
                    with open(file_path, "r", encoding="cp950") as f:
                        text = f.read()
                docs = [Document(page_content=text, metadata={"source": file.filename})]

            else:
                raise HTTPException(status_code=400, detail=f"不支援的檔案格式：{file.filename}")

            file_names.append(file.filename)
            documents.extend(docs)

        except Exception as e:
            raise HTTPException(status_code=500, detail=str(e))

    return documents, file_names
